"""Importación de animales desde una planilla Excel (.xlsx) o CSV.

Todo o nada: se valida el archivo completo y, si hay algún error, no se
importa ninguna fila. Así una carga inicial de cientos de animales no queda
a medias.
"""

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from livestock import services
from livestock.forms import FEMALE_CATEGORIES
from livestock.models import Animal, Breed, Farm, Lot, Paddock, Weighing, diio_validator

MAX_ROWS = 5000

# Columnas reconocidas (encabezado normalizado: minúsculas, sin tildes ni espacios).
COLUMNS = {
    "diio": {"diio", "crotal", "eid"},
    "tag": {"identificacion", "identificacionvisual", "autocrotal", "nombre", "tag"},
    "sex": {"sexo"},
    "category": {"categoria"},
    "breed": {"raza"},
    "birth_date": {"fechanacimiento", "nacimiento", "fechadenacimiento"},
    "farm": {"predio"},
    "paddock": {"potrero"},
    "lot": {"lote"},
    "weight": {"peso", "pesokg", "ultimopeso"},
    "weight_date": {"fechapeso", "fechadepeso", "fechadelpeso"},
    "mother": {"madre", "diiomadre"},
    "sire": {"padre", "toro"},
    "entry_date": {"fechaingreso", "ingreso", "fechadeingreso"},
}
REQUIRED = ["sex", "category", "farm"]
TEMPLATE_HEADER = [
    "DIIO", "Identificación", "Sexo", "Categoría", "Raza", "Fecha nacimiento", "Predio", "Potrero", "Lote",
    "Peso", "Fecha peso", "Madre", "Padre",
]
TEMPLATE_ROWS = [
    ["20345001", "", "Hembra", "Vaca", "Overo Colorado", "15-09-2020", "Fundo Los Robles", "Potrero 1",
     "Vacas de cría", "480", "01-10-2026", "", ""],
    ["", "T-104", "Macho", "Ternero", "Cruza", "10-08-2026", "Fundo Los Robles", "Potrero 1", "", "", "",
     "20345001", "ANG-123"],
]
SEXES = {"h": "F", "hembra": "F", "f": "F", "m": "M", "macho": "M"}


@dataclass
class ImportReport:
    errors: list = field(default_factory=list)
    created: int = 0
    rows: int = 0
    new_paddocks: set = field(default_factory=set)
    new_lots: set = field(default_factory=set)


def read_table(upload):
    """Filas (listas de textos) de un .xlsx o .csv."""
    data = upload.read()
    if upload.name.lower().endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        return [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    return services.csv_rows(data)


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime | date):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _choice(value, choices):
    """Valor de un TextChoices a partir de su etiqueta o valor, sin distinguir mayúsculas."""
    value = value.strip().lower()
    for v, label in choices:
        if value in (v.lower(), label.lower()):
            return v
    return None


def import_animals(organization, user, upload):
    report = ImportReport()
    table = [row for row in read_table(upload) if any(str(c).strip() for c in row)]
    if not table:
        report.errors.append("El archivo está vacío.")
        return report
    names = [services.norm_header(str(h)) for h in table[0]]
    cols = {key: next((i for i, n in enumerate(names) if n in options), None) for key, options in COLUMNS.items()}
    missing = [key for key in REQUIRED if cols[key] is None]
    if missing or (cols["diio"] is None and cols["tag"] is None):
        report.errors.append(
            "Faltan columnas obligatorias en la primera fila: Sexo, Categoría, Predio y DIIO o Identificación. "
            "Descarga la plantilla para ver el formato."
        )
        return report
    body = table[1:]
    if len(body) > MAX_ROWS:
        report.errors.append(f"El archivo tiene {len(body)} filas; el máximo por carga es {MAX_ROWS}.")
        return report
    report.rows = len(body)

    farms = {f.name.lower(): f for f in Farm.objects.for_org(organization)}
    breeds = {b.name.lower(): b for b in Breed.objects.all()}
    existing_diio = set(Animal.objects.for_org(organization).exclude(diio="").values_list("diio", flat=True))
    existing_tags = {
        t.lower() for t in Animal.objects.for_org(organization).exclude(tag="").values_list("tag", flat=True)
    }
    seen_diio, seen_tags, parsed = set(), set(), []
    today = date.today()

    for number, row in enumerate(body, start=2):
        def get(key, row=row):
            i = cols[key]
            return str(row[i]).strip() if i is not None and i < len(row) else ""

        errors = []
        diio, tag = get("diio").replace(" ", ""), get("tag")
        if not diio and not tag:
            errors.append("falta DIIO o identificación")
        if diio:
            try:
                diio_validator(diio)
            except ValidationError:
                errors.append(f"DIIO no válido ({diio})")
            if diio in existing_diio or diio in seen_diio:
                errors.append(f"DIIO {diio} repetido o ya registrado")
            seen_diio.add(diio)
        if tag:
            if tag.lower() in existing_tags or tag.lower() in seen_tags:
                errors.append(f"identificación {tag} repetida o ya registrada")
            seen_tags.add(tag.lower())
        sex = SEXES.get(get("sex").lower())
        if not sex:
            errors.append(f"sexo no válido ({get('sex')!r}); usa Hembra o Macho")
        category = _choice(get("category"), Animal.Category.choices)
        if not category:
            errors.append(f"categoría no válida ({get('category')!r})")
        elif sex and (category in FEMALE_CATEGORIES) != (sex == "F"):
            errors.append("la categoría no corresponde al sexo")
        breed = None
        if get("breed"):
            breed = breeds.get(get("breed").lower())
            if breed is None:
                errors.append(f"raza desconocida ({get('breed')!r})")
        farm = farms.get(get("farm").lower())
        if farm is None:
            errors.append(f"predio {get('farm')!r} no existe; créalo primero en Predios")
        dates = {}
        for key in ("birth_date", "weight_date", "entry_date"):
            if get(key):
                dates[key] = services.parse_date(get(key))
                if dates[key] is None or dates[key] > today:
                    errors.append(f"fecha no válida en {key.replace('_', ' ')} ({get(key)!r})")
        if dates.get("birth_date") and dates.get("entry_date") and dates["entry_date"] < dates["birth_date"]:
            errors.append("la fecha de ingreso es anterior al nacimiento")
        weight = None
        if get("weight"):
            weight = services.parse_weight(get("weight"))
            if weight is None:
                errors.append(f"peso no válido ({get('weight')!r})")
        if errors:
            report.errors.append(f"Fila {number}: " + "; ".join(errors) + ".")
            continue
        parsed.append({
            "diio": diio, "tag": tag, "sex": sex, "category": category, "breed": breed, "farm": farm,
            "paddock": get("paddock"), "lot": get("lot"), "weight": weight, "mother": get("mother"),
            "sire": get("sire")[:60], **dates,
        })

    # Madres: deben existir en el sistema o venir en el mismo archivo.
    mothers_db = {
        a.diio: a for a in Animal.objects.for_org(organization).filter(sex="F").exclude(diio="")
    }
    file_females = {p["diio"] for p in parsed if p["diio"] and p["sex"] == "F"}
    for p in parsed:
        if p["mother"] and p["mother"] not in mothers_db and p["mother"] not in file_females:
            report.errors.append(f"Madre {p['mother']} no encontrada (debe ser una hembra con ese DIIO).")
    if report.errors:
        return report

    with transaction.atomic():
        created = {}
        for p in parsed:
            paddock = lot = None
            if p["paddock"]:
                paddock, new = Paddock.objects.get_or_create(
                    organization=organization, farm=p["farm"], name=p["paddock"]
                )
                if new:
                    report.new_paddocks.add(str(paddock))
            if p["lot"]:
                lot, new = Lot.objects.get_or_create(organization=organization, name=p["lot"])
                if new:
                    report.new_lots.add(lot.name)
            animal = Animal(
                organization=organization, diio=p["diio"], tag=p["tag"], sex=p["sex"], category=p["category"],
                breed=p["breed"], birth_date=p.get("birth_date"), sire=p["sire"], origin=Animal.Origin.INITIAL,
                entry_date=p.get("entry_date") or today, farm=p["farm"], paddock=paddock, lot=lot,
            )
            services.register_animal(animal, user)
            if p["weight"]:
                Weighing.objects.create(
                    organization=organization, animal=animal, date=p.get("weight_date") or today,
                    weight_kg=p["weight"], notes="Importado", created_by=user,
                )
            created[p["diio"] or id(p)] = (animal, p["mother"])
        for animal, mother in created.values():
            if mother:
                animal.mother = mothers_db.get(mother) or created[mother][0]
                animal.save(update_fields=["mother"])
        report.created = len(created)
    return report
